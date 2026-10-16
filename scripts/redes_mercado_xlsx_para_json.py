"""
Planilha da diretoria (MAIORES POR SEGMENTO <ano>.xlsx, aba BASE <ano>)
  -> database/data/visao-diretor/redes-mercado-<ano>.json

    python -I scripts/redes_mercado_xlsx_para_json.py <planilha.xlsx> <saida.json> [--aba "BASE 2026"]

Consumido pela página Maiores redes da Visão Diretor (docs/visao-diretor.md, "Maiores
redes"). Depois de gerado, o JSON é a fonte: a planilha não é mais lida pelo sistema.

Regras de limpeza (decisões do Tony, 2026-10-16):
- Parêntese com o DONO ("POSTOS SHELL (RAÍZEN)") ou com DESCRIÇÃO ("SMART FIT (ACADEMIAS)")
  some: fica só a empresa. O dono não entra em nada, nem na vinculação.
- Bandeiras somadas numa linha ("CARREFOUR BRASIL (CARREFOUR / ATACADÃO / SAM'S CLUB)")
  viram uma linha por bandeira. O número de lojas da planilha é a SOMA, então cada bandeira
  sai com o número pesquisado em LOJAS, ou sem número (qualidade "sem_numero") quando não
  havia dado público — a soma original fica em `lojas_linha_original`.
- Variante com o mesmo começo de nome ("LEROY MERLIN (+ LEROY MERLIN EXPRESS)") não é
  bandeira: o casamento por começo de nome já a alcança.

As decisões linha a linha moram em DECISOES (chave = RANK da planilha). Linha que não está
lá segue a regra padrão: tira o parêntese, mantém o resto. Mapeamento de coluna é sempre
pelo NOME do cabeçalho, nunca pela posição.
"""

import argparse
import json
import re
import sys

from openpyxl import load_workbook

COLUNAS = {
    "posicao": "RANK",
    "categoria": "SEGMENTO",
    "empresa": "EMPRESA",
    "uf": "UF",
    "lojas": "LOJAS",
    "data_base": "DATA BASE",
    "status": "STATUS",
    "fonte": "FONTE / OBS",
}

# Rótulo da planilha -> chave da categoria (App\Services\VisaoDiretor\Redes\CategoriasDeRede).
CATEGORIAS = {
    "DROGARIAS": "drogarias",
    "ESTACIONAMENTOS": "estacionamentos",
    "POSTOS DE GASOLINA": "postos",
    "MATERIAL DE CONST.": "material_construcao",
    "ALIMENTAÇÃO": "alimentacao",
    "REDE DE LOJAS": "rede_lojas",
    "SUPERMERCADOS / ATACAREJO": "supermercados",
    "CONVENIÊNCIA / AUTOSSERVIÇO": "conveniencia",
    "BELEZA / SAÚDE / BEM-ESTAR": "beleza_saude",
    "ÓTICAS": "oticas",
    "SERVIÇOS / FRANQUIAS": "servicos_franquias",
}

QUALIDADE = {
    "Atualizado": "atualizado",
    "Estimativa": "estimativa",
    "Planilha anterior": "planilha_anterior",
}


def b(nome, apelidos=None):
    """Uma bandeira nascida da quebra de uma linha somada: lojas a pesquisar."""
    return {"nome": nome, "apelidos": apelidos or [], "bandeira": True}


def n(nome, apelidos=None):
    """A linha continua uma só; só o nome (e os apelidos de busca) mudam."""
    return {"nome": nome, "apelidos": apelidos or [], "bandeira": False}


def pular(motivo):
    return {"pular": motivo}


# RANK -> decisão. Lista = as linhas em que a da planilha vira; pular = sai, com o motivo.
DECISOES = {
    # --- dono ou descrição: fica a empresa (alguns ganham apelido de busca)
    1: [n("POSTOS PETROBRAS", ["POSTO PETROBRAS"])],
    2: [n("POSTOS SHELL", ["POSTO SHELL"])],
    3: [n("POSTOS IPIRANGA", ["POSTO IPIRANGA"])],
    5: [n("O BOTICÁRIO", ["BOTICARIO"])],
    6: [n("RAIA DROGASIL", ["DROGASIL", "DROGA RAIA", "DROGARIA RAIA"])],
    13: [n("FARMÁCIAS ASSOCIADAS")],
    16: [n("LUBRAX+", ["LUBRAX"])],
    21: [n("AM/PM", ["AM PM", "AMPM"])],
    41: [n("JET OIL")],
    111: [n("NOSSO BAR")],
    184: pular("Mesma rede do DIA BRASIL (linha 187)."),
    187: [n("DIA", ["DIA%", "DIA MARKET", "DIA BRASIL", "SUPERMERCADO DIA"])],
    197: pular("BOB'S repetido (já está em Alimentação, linha 38): conferir se é outra rede."),
    307: [n("MINUTO PÃO DE AÇÚCAR")],
    326: [n("PMP DROGARIA", ["PRECO MAIS POPULAR"])],
    338: [n("NOVELTY MODAS", ["NOVELTY"]), b("7 BELLO")],
    390: [n("STOK CENTER")],
    511: [n("SUPERMERCADO LÍDER")],
    554: pular("Mesma rede da linha 396 (CARVALHO, PI)."),
    # --- a mesma rede escrita duas vezes na planilha: fica a de cima do ranking
    284: pular("Mesma rede da linha 186 (BACIO DI LATTE)."),
    334: pular("Mesma rede da linha 198 (LOJAS TAQI), em outra categoria."),
    352: pular("Mesma rede da linha 351 (MILIUM TEM DE TUDO)."),
    434: pular("Mesma rede da linha 400 (SANTA LUCIA, ES)."),
    518: pular("Mesma rede da linha 479 (COOP)."),
    522: pular("Mesma rede da linha 478 (FARMÁCIA VALE VERDE)."),
    560: pular("Mesma rede da linha 469 (BARBOSA SUPERMERCADOS)."),
    614: pular("Mesma rede da linha 606 (SUPERMERCADOS MUNDIAL)."),
    677: pular("Mesma rede da linha 643 (CASA RENA), em outra categoria."),
    694: pular("Mesma rede da linha 535 (SUPERMERCADOS GUANABARA)."),
    713: pular("Mesma empresa da linha 227 (C VALE), em outra categoria."),
    662: [n("COSTA ATACADÃO")],
    720: [n("OJO")],
    # --- bandeiras somadas: uma linha por bandeira
    4: [b("CACAU SHOW"), b("BRIGADERIA")],
    10: [b("AREZZO"), b("HERING"), b("RESERVA"), b("FARM")],
    11: [b("OXXO")],  # LOJAS SELECT = SHELL SELECT, que já é a linha 34
    14: [b("FARMÁCIAS PAGUE MENOS", ["PAGUE MENOS"]), b("EXTRAFARMA")],
    17: [b("ÓTICAS CAROL", ["OTICA CAROL"]), b("GRANDVISION"), b("SUNGLASS HUT")],
    18: [b("DROGARIA SÃO PAULO"), b("DROGARIAS PACHECO", ["DROGARIA PACHECO"])],
    27: [b("KOPENHAGEN"), b("BRASIL CACAU")],
    30: [b("CHILLI BEANS"), b("ÓTICAS CHILLI BEANS"), b("ECO CHILLI")],
    35: [b("ÓTICAS DINIZ", ["OTICA DINIZ"]), b("DNZ EYEWEAR", ["DNZ"])],
    40: [b("NATURA"), b("AVON")],
    43: [b("CASAS BAHIA"), b("PONTO FRIO", ["PONTO"])],
    49: [b("CARREFOUR"), b("ATACADÃO"), b("SAM'S CLUB", ["SAMS CLUB"])],
    63: [b("SPOLETO"), b("CHINA IN BOX"), b("CASA DO PÃO DE QUEIJO")],
    65: [b("PÃO DE AÇÚCAR"), b("EXTRA MERCADO", ["MERCADO EXTRA"]), b("MINI EXTRA")],  # MINUTO é a linha 307
    67: [b("LOJAS RENNER", ["RENNER"]), b("YOUCOM"), b("CAMICADO"), b("ASHUA")],
    70: [b("DROGARIA TOTAL"), b("TOTAL POPULAR")],
    76: [b("DROGARIA CATARINENSE"), b("PREÇO POPULAR", ["FARMACIA PRECO POPULAR"])],
    107: [b("BRETAS"), b("GBARBOSA", ["G BARBOSA"]), b("GIGA ATACADO", ["GIGA"]), b("PREZUNIC")],
    108: [b("LOJAS RIACHUELO", ["RIACHUELO"]), b("CARTER'S", ["CARTERS"]), b("CASA RIACHUELO")],
    120: [b("GRILETTO"), b("MONTANA GRILL")],
    131: [b("FRANGO ASSADO"), b("PIZZA HUT"), b("KFC")],
    134: [b("MADERO"), b("JERONIMO")],
    135: [b("INDIGO"), b("PAREBEM")],
    137: [b("LOJAS GRAZZIOTIN", ["GRAZZIOTIN"]), b("TOTTAL"), b("PORMENOS")],
    139: [b("CASA & VÍDEO", ["CASA E VIDEO"]), b("LE BISCUIT")],
    156: [b("MIX ATACAREJO", ["MIX MATEUS"]), b("MATEUS SUPERMERCADOS", ["MATEUS"])],
    158: [b("HABIB'S", ["HABIBS"]), b("RAGAZZO")],
    159: [b("DROGASMIL"), b("DROGARIAS TAMOIO", ["TAMOIO"]), b("DROGARIA ROSÁRIO", ["ROSARIO"])],
    161: [b("RI HAPPY"), b("PB KIDS")],
    167: [b("CENTAURO"), b("NIKE STORE", ["NIKE"])],
    177: [b("WORLD TENNIS"), b("FUTEBOL SUPER STORE")],
    182: [b("VIA VENETO"), b("BROOKSFIELD")],
    191: [b("AUTHENTIC FEET"), b("ARTWALK")],
    198: [b("LOJAS TAQI", ["TAQI"]), b("IPLACE")],
    206: [b("OUTBACK"), b("ABBRACCIO", ["ABRACCIO"]), b("AUSSIE GRILL")],
    219: [b("LE LIS", ["LE LIS BLANC"]), b("JOHN JOHN"), b("DUDALINA")],
    232: [b("EPA SUPERMERCADOS", ["EPA"]), b("MINEIRÃO")],
    233: [b("AMIGÃO SUPERMERCADOS", ["AMIGAO"]), b("BOA SUPERMERCADOS")],
    238: [b("DI SANTINNI"), b("CAPODARTE")],
    261: [b("ASFAR"), b("DESCONTO FÁCIL")],
    268: [b("COLCCI"), b("FORUM")],
    283: [b("SUPER MUFFATO", ["MUFFATO"]), b("MAX ATACADISTA")],
    285: [b("LOJAS POMPÉIA", ["POMPEIA"]), b("GANG")],
    294: [b("COMPER"), b("FORT ATACADISTA")],
    310: [b("CARREFOUR EXPRESS"), b("CARREFOUR BAIRRO")],
    314: [b("SANTO REMÉDIO"), b("FLEXFARMA"), b("FARMABEM")],
    319: [b("MART MINAS"), b("DOM ATACADISTA")],
    322: [n("TATU BOLA")],
    325: [b("KOCH"), b("KOMPRÃO")],
    327: [b("SHOULDER"), b("ORIBA")],
    331: [n("BAHAMAS", ["BAHAMAS MIX", "BAHAMAS SUPERMERCADOS"])],
    336: [b("SAVEGNAGO"), b("PAULISTÃO")],
    346: [b("FAST SHOP"), b("A2YOU")],
    356: [b("LOJAS GUAIBIM", ["GUAIBIM"]), b("CASA + FÁCIL", ["CASA MAIS FACIL"])],
    367: [b("TOK&STOK", ["TOK STOK", "TOKSTOK"]), b("MOBLY")],
    412: [b("RISSUL"), b("UNIDASUL")],
    419: [b("TELHANORTE"), b("TUMELERO")],
    421: [b("ZAFFARI"), b("BOURBON")],
    425: [b("SUPER NOSSO"), b("APOIO MINEIRO")],
    436: [b("SODIMAC"), b("DICICO")],
    473: [b("SUPER BEAL", ["BEAL"]), b("FESTVAL")],
    497: [b("SUPERMERCADO IMPERATRIZ", ["IMPERATRIZ"]), b("BRASIL ATACADISTA")],
    507: [b("LOPES SUPERMERCADOS", ["LOPES"]), b("RAIZ SUPERMERCADOS")],
    508: [b("ST. MARCHÉ", ["ST MARCHE", "SAINT MARCHE"]), b("EMPÓRIO SANTA MARIA")],
    536: [b("PERUZZO SUPERMERCADOS", ["PERUZZO"]), b("ECOMIX")],
    537: [b("SUPERMERCADOS SÃO VICENTE"), b("CAVICCHIOLLI"), b("ARENA ATACADO")],
    556: [b("PASSARELA SUPERMERCADOS", ["PASSARELA"]), b("LANGUIRU")],
    558: [b("SUPERBOM"), b("COSTAZUL"), b("BARCELOS ATACADISTA")],
    608: [b("NORDESTÃO"), b("SUPERFÁCIL"), b("NRD")],
    642: [b("MIZUNO"), b("OLYMPIKUS"), b("UNDER ARMOUR")],
    652: [b("SUPERMERCADOS CUBATÃO"), b("SPASSO SABORES"), b("PRATKO")],
    751: [b("LEVEROS"), b("PROFIZ"), b("REFRIGÁS")],
    757: [b("SUPER FORMOSA"), b("FORMOSA MIX"), b("FORMOSA MAGAZINE")],
}

# Lojas pesquisadas (out/2026) das bandeiras quebradas, e as correções de linha inteira.
# (RANK, nome) -> (lojas, data_base, qualidade, fonte). "estimativa" = número aproximado
# ou derivado (resto da soma da linha); bandeira sem entrada aqui fica "sem_numero".
MO = "mercadoeconsumo.com.br"
GN = "gironews.com"
LOJAS = {
    (49, "CARREFOUR"): (105, "jun/26", "pesquisado", "InvestNews: 105 hipermercados no fim de junho/2026"),
    (49, "ATACADÃO"): (387, "jun/26", "pesquisado", "InvestNews: 387 lojas cash-and-carry (jun/26)"),
    (49, "SAM'S CLUB"): (58, "jun/26", "pesquisado", "InvestNews: 58 Sam's Club (jun/26)"),
    (310, "CARREFOUR EXPRESS"): (142, "jun/26", "pesquisado", "InvestNews: 142 lojas de conveniência (jun/26)"),
    (310, "CARREFOUR BAIRRO"): (5, "set/26", "estimativa", "26 em jun/26; 21 vendidas ao Grupo Pereira em set/26 (" + MO + ")"),
    (65, "PÃO DE AÇÚCAR"): (187, "mar/26", "pesquisado", "Diário do Pará (mar/26): 187 Pão de Açúcar"),
    (65, "EXTRA MERCADO"): (164, "mar/26", "pesquisado", "Diário do Pará (mar/26): 164 Extra Mercado"),
    (65, "MINI EXTRA"): (155, "mar/26", "pesquisado", "Diário do Pará (mar/26): 155 Mini Extra"),
    (307, "MINUTO PÃO DE AÇÚCAR"): (221, "mar/26", "pesquisado", "Diário do Pará (mar/26): 221 Minuto Pão de Açúcar"),
    (107, "GBARBOSA"): (150, "fev/25", "estimativa", MO + " (fev/25): mais de 150 lojas"),
    (107, "GIGA ATACADO"): (14, "2026", "pesquisado", GN + ": 14 unidades no estado de SP"),
    (107, "BRETAS"): (26, "fev/25", "pesquisado", MO + " (fev/25): 26 lojas em Goiás após a venda de MG"),
    (107, "PREZUNIC"): (33, "fev/25", "pesquisado", MO + " (fev/25): 33 lojas"),
    (67, "LOJAS RENNER"): (442, "2T26", "pesquisado", MO + " (17/08/2026): 442 Renner"),
    (67, "YOUCOM"): (161, "2T26", "pesquisado", MO + " (17/08/2026): 161 Youcom"),
    (67, "CAMICADO"): (100, "2T26", "pesquisado", MO + " (17/08/2026): 100 Camicado"),
    (67, "ASHUA"): (16, "2T26", "pesquisado", MO + " (17/08/2026): 16 Ashua"),
    (14, "EXTRAFARMA"): (348, "dez/24", "estimativa", "Balanço 4T24: 348 Extrafarma, antes de sair de 6 estados em 2025"),
    (14, "FARMÁCIAS PAGUE MENOS"): (1341, "dez/25", "estimativa", "Resto da linha (1.689) menos Extrafarma"),
    (18, "DROGARIA SÃO PAULO"): (1030, "jan/25", "estimativa", "DPSP 1.613 lojas; proporção de 2022 (~64% Drogaria São Paulo)"),
    (18, "DROGARIAS PACHECO"): (583, "jan/25", "estimativa", "DPSP 1.613 lojas; resto da proporção de 2022"),
    (4, "CACAU SHOW"): (4781, "abr/26", "pesquisado", MO + " (05/04/2026): 4.781 lojas"),
    (4, "BRIGADERIA"): (8, "jul/24", "estimativa", MO + " (16/07/2024): 8 lojas"),
    (10, "HERING"): (700, "set/25", "pesquisado", MO + " (04/09/2025): cerca de 700 lojas, 613 franqueadas"),
    (10, "RESERVA"): (110, "2020", "estimativa", "Aquisição pela Arezzo (2020): 110 lojas"),
    (17, "ÓTICAS CAROL"): (1400, "2026", "pesquisado", "EssilorLuxottica, ABF Expo 2026: mais de 1.400 Óticas Carol"),
    (27, "KOPENHAGEN"): (850, "mai/25", "pesquisado", "IstoÉ Dinheiro (25/05/2025): mais de 850 pontos de venda"),
    (27, "BRASIL CACAU"): (500, "mai/25", "pesquisado", "IstoÉ Dinheiro (25/05/2025): mais de 500 lojas"),
    (43, "PONTO FRIO"): (119, "jul/25", "estimativa", MO + " (15/07/2025): 119 lojas em repaginação"),
    (43, "CASAS BAHIA"): (923, "dez/25", "estimativa", "Resto da linha (1.042) menos Ponto Frio"),
    (63, "SPOLETO"): (400, "2025", "pesquisado", GN + ": abertura do 400º restaurante em 2025"),
    (63, "CHINA IN BOX"): (150, "mar/24", "estimativa", MO + " (26/03/2024): mais de 150 lojas"),
    (30, "ÓTICAS CHILLI BEANS"): (300, "mai/24", "estimativa", MO + " (09/05/2024): 300 óticas"),
    (131, "KFC"): (231, "2025", "pesquisado", "InfoMoney (2025): 231 unidades"),
    (131, "FRANGO ASSADO"): (23, "2025", "pesquisado", "InfoMoney (2025): 23ª unidade aberta"),
    (131, "PIZZA HUT"): (90, "s/d", "estimativa", "Tiendeo: cerca de 90 restaurantes (sem data)"),
    (156, "MIX ATACAREJO"): (98, "3T25", "pesquisado", "Banda B (3T25): 98 Mix Mateus"),
    (156, "MATEUS SUPERMERCADOS"): (204, "dez/25", "estimativa", "Resto da linha (302) menos Mix Mateus"),
    (159, "DROGARIA ROSÁRIO"): (65, "2023", "estimativa", "Release 3T23: 65 Drogaria Rosário"),
    (134, "MADERO"): (252, "4T25", "pesquisado", GN + " (4T25): 108 Steak House + 144 Burger"),
    (134, "JERONIMO"): (98, "4T25", "pesquisado", GN + " (4T25): 98 Jeronimo"),
    (108, "CARTER'S"): (81, "set/25", "pesquisado", MO + " (25/09/2025): 81 lojas"),
    (70, "DROGARIA TOTAL"): (585, "dez/23", "estimativa", GN + " (fim de 2023): 585 Drogaria Total"),
    (70, "TOTAL POPULAR"): (33, "dez/23", "estimativa", GN + " (fim de 2023): 33 Total Popular"),
    (120, "GRILETTO"): (150, "2019", "estimativa", "Portal do Franchising (2019): mais de 150 lojas"),
    (120, "MONTANA GRILL"): (245, "2026", "estimativa", "Resto da linha (395) menos Griletto"),
    (135, "INDIGO"): (354, "2025", "estimativa", "NeoFeed: 354 operações, já somando a PareBem (incorporada em 2022)"),
    (139, "CASA & VÍDEO"): (260, "2023", "estimativa", "BP Money (2023): cerca de 260 lojas"),
    (139, "LE BISCUIT"): (139, "2023", "estimativa", "BP Money (2023): 139 unidades"),
    (167, "NIKE STORE"): (53, "set/25", "pesquisado", "Bloomberg Línea (set/25): 53 lojas"),
    (167, "CENTAURO"): (230, "nov/25", "estimativa", "BP Money (nov/25): presente em cerca de 230 shoppings"),
    (177, "WORLD TENNIS"): (260, "s/d", "estimativa", "Tiendeo: 260 lojas (sem data)"),
    (191, "ARTWALK"): (73, "2023", "estimativa", "Bem Paraná (2023): 73 lojas"),
    (191, "AUTHENTIC FEET"): (100, "s/d", "estimativa", "Mais de 100 lojas físicas (sem data)"),
    (198, "LOJAS TAQI"): (22, "nov/25", "estimativa", "Jornal do Comércio (nov/25): 29 das 51 lojas fechando"),
    (198, "IPLACE"): (137, "s/d", "estimativa", "Material do grupo Herval (sem data): 137 lojas"),
    (206, "OUTBACK"): (174, "s/d", "estimativa", "InvestNews (sem data): 174 Outback"),
    (206, "ABBRACCIO"): (17, "s/d", "estimativa", "InvestNews (sem data): 17 Abbraccio"),
    (206, "AUSSIE GRILL"): (2, "s/d", "estimativa", "InvestNews (sem data): 2 Aussie"),
    (219, "LE LIS"): (108, "s/d", "estimativa", "Wikipedia (sem data): 108 Le Lis"),
    (219, "DUDALINA"): (58, "2025", "pesquisado", "InvestNews (2025): 37 próprias + 21 franquias"),
    (219, "JOHN JOHN"): (55, "s/d", "estimativa", "Wikipedia (sem data): 55 John John"),
    (233, "AMIGÃO SUPERMERCADOS"): (66, "2025", "pesquisado", GN + ": 66 unidades em SP, PR e MS"),
    (238, "DI SANTINNI"): (115, "2023", "estimativa", "Mais de 115 lojas (2023)"),
    (238, "CAPODARTE"): (80, "2025", "estimativa", "80 lojas (12 próprias e 68 franquias)"),
    (268, "COLCCI"): (102, "s/d", "estimativa", "IstoÉ Dinheiro (sem data): 102 Colcci"),
    (268, "FORUM"): (45, "s/d", "estimativa", "IstoÉ Dinheiro (sem data): 45 Forum"),
    (285, "LOJAS POMPÉIA"): (83, "2022", "estimativa", "Jornal do Comércio (2022): 83 lojas"),
    (285, "GANG"): (50, "fev/23", "estimativa", "Jornal do Comércio (fev/23): mais de 50 filiais"),
    (294, "FORT ATACADISTA"): (77, "jul/26", "pesquisado", GN + " (01/07/2026): 77 Fort Atacadista"),
    (294, "COMPER"): (40, "jul/26", "pesquisado", GN + " (01/07/2026): 40 Comper, antes das 21 lojas do Carrefour"),
}

PARENTESE = re.compile(r"\s*\(.*?\)\s*")


def nome_padrao(empresa: str) -> str:
    return re.sub(r"\s+", " ", PARENTESE.sub(" ", empresa)).strip()


def ler(planilha: str, aba: str):
    wb = load_workbook(planilha, read_only=True, data_only=True)
    linhas = wb[aba].iter_rows(values_only=True)
    cabecalho = [str(c).strip() if c is not None else "" for c in next(linhas)]
    try:
        idx = {chave: cabecalho.index(col) for chave, col in COLUNAS.items()}
    except ValueError as e:
        sys.exit(f"Cabeçalho inesperado na aba {aba}: {e}")
    for linha in linhas:
        if linha[idx["posicao"]] is None:
            continue
        yield {chave: linha[i] for chave, i in idx.items()}


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("planilha")
    ap.add_argument("saida")
    ap.add_argument("--aba", default="BASE 2026")
    args = ap.parse_args()

    redes, puladas, vistas = [], [], set()
    usadas, usadas_lojas = set(), set()
    for linha in ler(args.planilha, args.aba):
        posicao = int(linha["posicao"])
        rotulo = str(linha["categoria"]).strip()
        if rotulo not in CATEGORIAS:
            sys.exit(f"Categoria desconhecida na linha {posicao}: {rotulo}")
        original = str(linha["empresa"]).strip()
        decisao = DECISOES.get(posicao)
        usadas.add(posicao)

        if isinstance(decisao, dict):
            puladas.append({"posicao": posicao, "linha_original": original, "motivo": decisao["pular"]})
            continue

        itens = decisao or [n(nome_padrao(original))]
        uf = str(linha["uf"]).strip() if linha["uf"] else None
        lojas = int(linha["lojas"]) if linha["lojas"] is not None else None

        for item in itens:
            chave = (CATEGORIAS[rotulo], item["nome"].upper())
            if chave in vistas:
                sys.exit(f"Rede repetida na mesma categoria: {chave} (linha {posicao})")
            vistas.add(chave)
            bandeira = item["bandeira"] and len(itens) > 1
            if bandeira:
                lj, base, qual, fonte = None, None, "sem_numero", None
            else:
                lj = lojas
                base = str(linha["data_base"]) if linha["data_base"] else None
                qual = QUALIDADE.get(str(linha["status"]).strip(), "estimativa")
                fonte = linha["fonte"]
            pesquisa = LOJAS.get((posicao, item["nome"]))
            if pesquisa:
                usadas_lojas.add((posicao, item["nome"]))
                lj, base, qual, fonte = pesquisa
            redes.append({
                "posicao": posicao,
                "categoria": CATEGORIAS[rotulo],
                "nome": item["nome"],
                "apelidos": item["apelidos"],
                "uf": None if uf in (None, "-", "") else uf,
                "lojas": lj,
                "data_base": base,
                "qualidade": qual,
                "fonte": fonte,
                "linha_original": original,
                "lojas_linha_original": lojas,
            })

    sobrando = sorted(set(DECISOES) - usadas)
    if sobrando:
        sys.exit(f"Decisões para linhas que não existem na planilha: {sobrando}")
    sem_rede = sorted(set(LOJAS) - usadas_lojas)
    if sem_rede:
        sys.exit(f"Lojas pesquisadas para redes que não existem: {sem_rede}")

    with open(args.saida, "w", encoding="utf-8") as f:
        json.dump({
            "fonte": f"scripts/redes_mercado_xlsx_para_json.py sobre a aba '{args.aba}'",
            "redes": redes,
            "puladas": puladas,
        }, f, ensure_ascii=False, indent=1)
        f.write("\n")

    sem_numero = sum(1 for r in redes if r["qualidade"] == "sem_numero")
    print(f"{len(redes)} redes, {len(puladas)} linhas puladas, {sem_numero} bandeiras sem número de lojas")


if __name__ == "__main__":
    main()
